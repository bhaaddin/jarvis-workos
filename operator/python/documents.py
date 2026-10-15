"""Fixed document operations. JSON stdin/stdout; no generated Python runs here."""
import csv
from decimal import Decimal
import hashlib
import io
import json
import os
from pathlib import Path
import sys
import zipfile
from xml.sax.saxutils import escape

MAX_BYTES = 20 * 1024 * 1024
MAX_TEXT = 48000


def check_archive(path):
    if path.suffix.lower() in (".docx", ".xlsx"):
        with zipfile.ZipFile(path) as archive:
            entries = archive.infolist()
            if len(entries) > 4000 or sum(item.file_size for item in entries) > 60 * 1024 * 1024:
                raise ValueError("Document archive exceeds extraction limits")


def read_document(path, start=1, limit=10):
    if not path.is_file() or path.stat().st_size > MAX_BYTES:
        raise ValueError("Choose a document no larger than 20 MiB")
    check_archive(path)
    kind = path.suffix.lower()
    start, limit = max(1, int(start)), min(300, max(1, int(limit)))
    result = {"name": path.name, "format": kind[1:], "sha256": hashlib.sha256(path.read_bytes()).hexdigest()}
    if kind == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(path, strict=False)
        if reader.is_encrypted:
            raise ValueError("This PDF is encrypted; provide an unlocked copy")
        pages = [{"page": index + 1, "text": (reader.pages[index].extract_text() or "")[:MAX_TEXT]}
                 for index in range(start - 1, min(len(reader.pages), start - 1 + min(limit, 30)))]
        result.update(pages=pages, total_pages=len(reader.pages), nextStart=start+len(pages) if start-1+len(pages)<len(reader.pages) else None,
                      needs_ocr=not any(page["text"].strip() for page in pages))
    elif kind == ".docx":
        from docx import Document
        document = Document(path)
        paragraphs = [paragraph.text for paragraph in document.paragraphs]
        tables = [[[cell.text[:2000] for cell in row.cells[:30]] for row in table.rows[:200]] for table in document.tables[:10]]
        result.update(text="\n".join(paragraphs)[:MAX_TEXT], tables=tables)
    elif kind == ".xlsx":
        from openpyxl import load_workbook
        book = load_workbook(path, read_only=True, data_only=False, keep_links=False)
        try:
            result["sheets"] = [{"name": sheet.title, "rows": [[normalize(cell) for cell in row]
                                 for row in sheet.iter_rows(min_row=start, max_row=min(sheet.max_row, start + min(200, limit) - 1), max_col=min(30, sheet.max_column), values_only=True)],
                                 "total_rows": sheet.max_row, "nextStart": start+min(200, limit) if start-1+min(200, limit)<sheet.max_row else None}
                                for sheet in list(book)[:10]]
        finally:
            book.close()
    elif kind == ".csv":
        with path.open(encoding="utf-8-sig", newline="") as stream:
            rows, more = [], False
            for index, row in enumerate(csv.reader(stream)):
                if index >= start - 1 + min(200, limit):
                    more = True
                    break
                if index >= start - 1:
                    rows.append([cell[:2000] for cell in row[:30]])
            result["rows"] = rows
            result["nextStart"] = start + len(rows) if more else None
    elif kind in (".txt", ".md", ".json", ".log"):
        text = path.read_text(encoding="utf-8-sig")
        lines = text.splitlines()
        selected = lines[start - 1:start - 1 + limit]
        result.update(text="\n".join(selected)[:MAX_TEXT], total_lines=len(lines), start=start,
                      nextStart=start+len(selected) if start-1+len(selected)<len(lines) else None)
    else:
        raise ValueError("Supported inputs: PDF, DOCX, XLSX, CSV, Markdown, text, JSON and text logs")
    if len(json.dumps(result, ensure_ascii=False)) > 180000:
        raise ValueError("Extraction is too large; request fewer pages/rows")
    return result


def normalize(value):
    if value is None or isinstance(value, (str, bool, int, float)):
        return value
    return str(value)


def safe_cell(value):
    # Text is data, never an Excel/DDE/external-link formula.
    return "'" + value if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@")) else value


def load_table(path, sheet=None):
    if not path.is_file() or path.stat().st_size > MAX_BYTES:
        raise ValueError('Choose a table no larger than20MiB')
    check_archive(path)
    book = None
    stream = None
    try:
        if path.suffix.lower() == '.csv':
            stream = path.open(encoding='utf-8-sig', newline='')
            sample = stream.read(8192)
            stream.seek(0)
            try:
                dialect = csv.Sniffer().sniff(sample, delimiters=',;\t')
            except csv.Error:
                dialect = csv.excel
            iterator = iter(csv.reader(stream, dialect))
        elif path.suffix.lower() == '.xlsx':
            from openpyxl import load_workbook
            book = load_workbook(path, read_only=True, data_only=True, keep_links=False)
            iterator = iter((book[sheet] if sheet else book.worksheets[0]).iter_rows(values_only=True))
        else:
            raise ValueError('Table analysis accepts CSV or XLSX only')
        columns = [str(value or '').strip() for value in next(iterator)]
        if not 1 <= len(columns) <= 30 or any(not value or len(value) > 200 for value in columns):
            raise ValueError('Use1–30 nonempty column headers')
        rows = []
        for row in iterator:
            if not any(value is not None and value != '' for value in row):
                continue
            if len(rows) >= 50000 or len(row) > len(columns):
                raise ValueError('The complete table exceeds its row/column bound; no partial total was computed')
            rows.append([None if value is None else str(value) for value in row] + [None] * (len(columns) - len(row)))
        return {'columns': columns, 'rows': rows, 'sha256': hashlib.sha256(path.read_bytes()).hexdigest()}
    finally:
        if book:
            book.close()
        if stream:
            stream.close()


def create_document(path, spec):
    title = str(spec.get("title", "JARVIS document"))[:300]
    content = str(spec.get("content", ""))[:60000]
    columns = [str(value)[:200] for value in spec.get("columns", [])][:30]
    rows = spec.get("rows", [])
    if len(rows) > 5000 or any(not isinstance(row, list) or len(row) > 30 for row in rows):
        raise ValueError("Use no more than 5000 rows and 30 columns")
    totals = {}
    for column in spec.get("totalColumns", []):
        index = int(column)
        if index < 0 or index >= len(columns):
            raise ValueError("Total-column index is outside the table")
        values = [row[index] for row in rows if index < len(row) and row[index] is not None]
        if any(isinstance(value, bool) or not isinstance(value, (int, float)) for value in values):
            raise ValueError("Total columns must contain numbers")
        totals[str(index)] = str(sum((Decimal(str(value)) for value in values), Decimal(0)))
    if totals:
        total = [None] * len(columns)
        for index, value in totals.items():
            total[int(index)] = float(Decimal(value))
        if "0" not in totals:
            total[0] = "Total"
        rows = [*rows, total]
    kind = path.suffix.lower()
    if kind in (".txt", ".md", ".json"):
        if kind == ".json":
            content = json.dumps(json.loads(content), indent=2, ensure_ascii=False)
        path.write_text(content if kind != ".md" else f"# {title}\n\n{content}\n", encoding="utf-8")
    elif kind == ".csv":
        with path.open("w", encoding="utf-8-sig", newline="") as stream:
            writer = csv.writer(stream)
            if columns:
                writer.writerow([safe_cell(value) for value in columns])
            writer.writerows([[safe_cell(value) for value in row] for row in rows])
    elif kind == ".xlsx":
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        book = Workbook()
        sheet = book.active
        sheet.title = "Report"
        if columns:
            sheet.append([safe_cell(value) for value in columns])
            for cell in sheet[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="174E63")
            sheet.freeze_panes = "A2"
        for row in rows:
            sheet.append([safe_cell(value) for value in row])
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = min(50, max(14, max(len(str(cell.value or "")) for cell in column) + 2))
        if content:
            notes = book.create_sheet("Notes")
            for line in content.splitlines():
                notes.append([safe_cell(line)])
        book.save(path)
        book.close()
    elif kind == ".docx":
        from docx import Document
        document = Document()
        document.add_heading(title, 0)
        for line in content.splitlines():
            document.add_paragraph(line)
        if columns:
            table = document.add_table(rows=1, cols=len(columns))
            table.style = "Light Shading Accent 1"
            for cell, value in zip(table.rows[0].cells, columns):
                cell.text = value
            for row in rows:
                for cell, value in zip(table.add_row().cells, row):
                    cell.text = str(value if value is not None else "")
        document.save(path)
    elif kind == ".pdf":
        from reportlab.lib import colors
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.lib.pagesizes import A4
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        font = Path(os.environ.get("WINDIR", "C:/Windows")) / "Fonts/arial.ttf"
        if font.is_file():
            pdfmetrics.registerFont(TTFont("JarvisText", str(font)))
        styles = getSampleStyleSheet()
        for style in styles.byName.values():
            style.fontName = "JarvisText" if font.is_file() else "Helvetica"
        story = [Paragraph(escape(title), styles["Title"]), Spacer(1, 12)]
        story.extend(Paragraph(escape(line) or "&nbsp;", styles["BodyText"]) for line in content.splitlines())
        if columns:
            table = Table([[Paragraph(escape(str(value if value is not None else "")), styles["BodyText"]) for value in row]
                           for row in [columns, *rows]], repeatRows=1)
            table.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#dceaf0")),
                                      ("VALIGN", (0, 0), (-1, -1), "TOP"), ("GRID", (0, 0), (-1, -1), .25, colors.lightgrey)]))
            story.extend([Spacer(1, 12), table])
        SimpleDocTemplate(str(path), pagesize=A4, title=title).build(story)
    else:
        raise ValueError("Supported outputs: PDF, DOCX, XLSX, CSV, Markdown, text and JSON")
    if path.stat().st_size > MAX_BYTES:
        path.unlink()
        raise ValueError("Generated document exceeded the size limit")
    return {"bytes": path.stat().st_size, "sha256": hashlib.sha256(path.read_bytes()).hexdigest(), "totals": totals,
            "tableRowsWritten": len(rows), "columnsWritten": columns, "verifiedPreview": read_document(path, 1, 2 if kind == ".pdf" else 20)}


if __name__ == "__main__":
    try:
        sys.stdin.reconfigure(encoding="utf-8")
        sys.stdout.reconfigure(encoding="utf-8")
        request = json.loads(sys.stdin.read(1024 * 1024))
        path = Path(request["path"])
        result = load_table(path, request.get('sheet')) if request['operation'] == 'table' else read_document(path, request.get("start", 1), request.get("limit", 10)) if request["operation"] == "read" else create_document(path, request["spec"])
        print(json.dumps({"ok": True, "result": result}, ensure_ascii=False))
    except Exception as error:
        print(json.dumps({"ok": False, "error": str(error)[:600]}, ensure_ascii=False))
        sys.exit(1)
